import logging
from abc import ABC
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt"}


class DocumentParser(ABC):
    @staticmethod
    def parse(file_path: Path) -> str:
        raise NotImplementedError


class PDFParser(DocumentParser):
    @staticmethod
    def parse(file_path: Path) -> str:
        if not file_path.exists():
            raise FileNotFoundError(f"PDF file not found: {file_path}")

        text = []
        try:
            reader = PdfReader(file_path)
            for page in reader.pages:
                page_text = page.extract_text()
                if not page_text:
                    continue
                text.append(page_text)

            logging.info(f"Parsed {len(reader.pages)} pages from {file_path.name}")
        except Exception as e:
            logging.error(f"Error parsing PDF {file_path}: {e}")
            raise

        return "\n\n".join(text)


class DocxParser(DocumentParser):
    @staticmethod
    def parse(file_path: Path) -> str:
        if not file_path.exists():
            raise FileNotFoundError(f"DOCX file not found: {file_path}")

        text = []
        try:
            doc = DocxDocument(str(file_path))
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text.append(paragraph.text)

            # Also extract tables if present
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text for cell in row.cells)
                    if not row_text.strip():
                        continue
                    text.append(row_text)

            logging.info(f"Parsed {len(doc.paragraphs)} paragraphs from {file_path.name}")
        except Exception as e:
            logging.error(f"Error parsing DOCX {file_path}: {e}")
            raise

        return "\n\n".join(text)


class TxtParser(DocumentParser):
    @staticmethod
    def parse(file_path: Path) -> str:
        if not file_path.exists():
            raise FileNotFoundError(f"TXT file not found: {file_path}")

        try:
            with file_path.open("r", encoding="utf-8") as f:
                text = f.read()
            logging.info(f"Parsed TXT file {file_path.name}")
            return text
        except Exception as e:
            logging.error(f"Error parsing TXT {file_path}: {e}")
            raise


class ParserFactory:
    _parsers = {".pdf": PDFParser, ".docx": DocxParser, ".doc": DocxParser, ".txt": TxtParser}

    @classmethod
    def get_parser(cls, file_path: Path) -> DocumentParser:
        suffix = file_path.suffix.lower()
        parser_class = cls._parsers.get(suffix)
        if not parser_class:
            raise ValueError(f"Unsupported file type: {suffix}")
        return parser_class

    @classmethod
    def parse(cls, file_path: Path) -> str:
        parser = cls.get_parser(file_path)
        return parser.parse(file_path)
